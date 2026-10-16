"""Smart spreadsheet reader + column mapper for the product bulk import.

Pharmacies migrate from many systems (Odoo, ERPNext, Dolibarr, local Egyptian
pharmacy software, hand-made Excel sheets …) whose exports differ in:

* file format    — .xlsx, legacy .xls, HTML tables saved as .xls, CSV/TSV;
* text encoding  — UTF-8, UTF-16, Windows-1256 (Arabic), Latin-1;
* layout         — title/blank rows above the header, several sheets;
* header language — English, Arabic, French, Spanish, Turkish, German …;
* unit codes     — PAC/STR/TAB, Box/Strip, علبة/شريط …

This module turns any of those into rows keyed by the canonical header names
that ``inventory._parse_bulk_row`` already understands, and reports the
detected mapping so the UI can show it (and let the user correct it).
"""
from __future__ import annotations

import csv
import io
import re
import unicodedata
from datetime import date, datetime, timedelta
from html.parser import HTMLParser
from typing import Any, Optional

# ─── Canonical fields ──────────────────────────────────────────────────────
# key → (canonical header consumed by _parse_bulk_row, English label, Arabic label, required)
FIELDS: dict[str, tuple[str, str, str, bool]] = {
    "barcode":        ("barcode", "Item code", "كود الصنف", False),
    "intl_barcode":   ("international barcode", "International barcode", "الباركود الدولي", False),
    "name_en":        ("name_en", "Item name", "اسم الصنف", True),
    "name_ar":        ("name_ar", "Arabic name", "الاسم بالعربي", False),
    "unit":           ("unit", "Main unit (box)", "الوحدة الكبرى (علبة)", False),
    "sub_unit":       ("sub_unit", "Sub unit (strip)", "الوحدة الصغرى (شريط)", False),
    "pack_size":      ("pack_size", "Sub units per box", "عدد الوحدات الصغرى في العلبة", False),
    "stock":          ("stock", "Stock (boxes)", "الرصيد (علب)", False),
    "price":          ("price", "Selling price", "سعر البيع", False),
    "sub_price":      ("sub_price", "Sub-unit (strip) price", "سعر الوحدة الصغرى", False),
    "cost":           ("cost", "Cost price", "سعر التكلفة", False),
    "discount_pct":   ("discount %", "Discount / margin %", "نسبة الخصم / الربح", False),
    "category":       ("category", "Category", "التصنيف", False),
    "material_group": ("material_group", "Material group", "مجموعة الصنف", False),
    "origin":         ("origin_type", "Local / imported", "محلي / مستورد", False),
    "min_stock":      ("min_stock", "Min stock", "الحد الأدنى", False),
    "expiry":         ("expiry_date", "Expiry date", "تاريخ الصلاحية", False),
}

# Synonyms are matched after normalize_header(); Arabic is normalised the same way.
SYNONYMS: dict[str, list[str]] = {
    "barcode": [
        "code", "item code", "material code", "product code", "internal code", "internal reference",
        "reference", "ref", "sku", "item no", "item number", "item id", "product id", "default code",
        "article code", "code article", "reference interne", "codigo", "codigo articulo", "codigo producto",
        "stok kodu", "urun kodu", "artikelnummer", "artikel nr", "codice", "codigo do produto",
        "كود", "الكود", "كود الصنف", "رقم الصنف", "كود الماده", "كود المنتج", "الرقم", "كود داخلي", "رمز الصنف",
    ],
    "barcode_generic": [
        "barcode", "barcodes", "bar code", "ean", "ean13", "ean 13", "upc", "gtin", "code barre", "code barres",
        "code a barres", "codigo de barras", "barkod", "strichcode", "codice a barre",
        "باركود", "الباركود", "الباركود الدولي", "كود دولي",
    ],
    "intl_barcode": [
        "international barcode", "international_barcode", "intl barcode", "global barcode", "global code",
        "international code", "gtin barcode", "ean code",
    ],
    "name_en": [
        "name", "name_en", "item name", "material name", "product name", "product", "item", "description",
        "item description", "english name", "name english", "name en", "trade name", "drug name", "medicine",
        "label", "designation", "libelle", "nom", "nom du produit", "article", "nombre", "descripcion",
        "producto", "articulo", "urun adi", "urun", "stok adi", "malzeme adi", "bezeichnung", "artikel",
        "artikelname", "produktname", "nome", "descricao", "nome do produto",
        "الاسم", "اسم", "اسم الصنف", "الصنف", "اسم المنتج", "اسم الماده", "اسم الدواء", "البيان", "الوصف",
        "الاسم التجاري", "الاسم الانجليزي", "اسم انجليزي", "الاسم بالانجليزي",
    ],
    "name_ar": [
        "name_ar", "arabic name", "name arabic", "name ar", "name (arabic)", "arabic", "arabic description",
        "الاسم العربي", "اسم عربي", "الاسم بالعربي", "اسم الصنف عربي", "الاسم بالعربية",
    ],
    "unit": [
        "unit", "main unit", "big unit", "large unit", "material unit", "uom", "unit of measure", "stock uom",
        "sales unit", "package", "pack unit", "major unit", "unite", "unidad", "unidad de medida", "birim",
        "ana birim", "einheit", "unidade",
        "الوحده", "وحده", "الوحده الكبري", "الوحده الكبيره", "وحده كبري", "وحده البيع", "العبوه",
    ],
    "sub_unit": [
        "sub unit", "sub_unit", "subunit", "small unit", "small unit name", "minor unit", "inner unit",
        "child unit", "secondary unit", "unit classification", "small unit type", "sub uom",
        "sous unite", "unite secondaire", "subunidad", "unidad menor", "alt birim", "untereinheit",
        "الوحده الصغري", "وحده صغري", "الوحده الصغيره", "الوحده الفرعيه", "التجزئه", "وحده التجزئه",
    ],
    "pack_size": [
        "pack size", "pack_size", "subunit quantity", "sub unit quantity", "subunit qty", "sub unit qty",
        "small unit quantity per unit", "small unit qty per unit", "small unit per unit", "small unit quantity",
        "units per unit", "quantity per unit", "qty per unit", "number of small unit", "number of small units",
        "small units", "units per pack", "units per box", "units per 1 box", "units per big unit",
        "conversion factor", "conversion", "factor", "uom factor", "pieces per box", "strips per box",
        "content", "contenance", "qte par unite", "unidades por caja", "factor de conversion", "carpan",
        "inhalt", "umrechnungsfaktor",
        "عدد الوحدات الصغري", "عدد الوحدات", "عدد الشرائط", "عدد الاشرطه", "معامل التحويل", "التحويل",
        "عدد التجزئه", "محتوي العبوه", "الكميه في العلبه",
    ],
    "stock": [
        "stock", "quantity", "qty", "qty on hand", "quantity on hand", "on hand", "balance", "available",
        "available qty", "current stock", "opening stock", "stock qty", "inventory", "count", "soh",
        "quantite", "qte", "stock disponible", "cantidad", "existencia", "existencias", "inventario",
        "miktar", "stok", "stok miktari", "bestand", "menge", "lagerbestand", "quantidade", "estoque",
        "الكميه", "كميه", "الرصيد", "رصيد", "المخزون", "الرصيد الحالي", "الكميه المتاحه", "الموجود",
        "رصيد اول المده", "الكميه الحاليه",
    ],
    "price": [
        "price", "sales price", "selling price", "sale price", "retail price", "public price", "unit price",
        "list price", "material price", "standard selling rate", "selling rate", "mrp", "consumer price",
        "prix", "prix de vente", "prix public", "prix unitaire", "precio", "precio venta", "precio de venta",
        "pvp", "satis fiyati", "fiyat", "verkaufspreis", "preis", "preco", "preco de venda",
        "السعر", "سعر", "سعر البيع", "سعر الجمهور", "سعر المستهلك", "سعر العلبه", "البيع", "سعر الوحده",
    ],
    "sub_price": [
        "sub price", "sub_price", "sub unit price", "subunit price", "small unit price", "strip price",
        "piece price", "unit price small", "retail sub price", "prix detail",
        "سعر الشريط", "سعر الوحده الصغري", "سعر التجزئه", "سعر القطعه",
    ],
    "cost": [
        "cost", "cost price", "purchase price", "buying price", "buy price", "unit cost", "average cost",
        "standard price", "valuation rate", "last purchase rate", "pharmacist price", "supplier price",
        "prix d achat", "prix de revient", "cout", "precio compra", "precio de compra", "costo", "coste",
        "alis fiyati", "maliyet", "einkaufspreis", "kosten", "preco de custo", "custo",
        "التكلفه", "تكلفه", "سعر التكلفه", "سعر الشراء", "الشراء", "سعر الصيدلي", "متوسط التكلفه",
    ],
    "discount_pct": [
        "discount %", "discount_pct", "discount percent", "discount percentage", "discount", "profit margin %",
        "margin %", "profit %", "margin", "markup", "remise", "descuento", "iskonto", "rabatt",
        "نسبه الخصم", "الخصم", "خصم", "نسبه الربح", "هامش الربح",
    ],
    "category": [
        "category", "product category", "item group", "group", "family", "class", "classification", "type",
        "categorie", "famille", "categoria", "familia", "grupo", "kategori", "grup", "kategorie", "gruppe",
        "التصنيف", "تصنيف", "الفئه", "المجموعه", "القسم", "النوع", "الفصيله",
    ],
    "material_group": [
        "material group", "material_group", "materialgroup", "mg", "mat group", "product group",
        "material type", "material classification", "مجموعه الماده", "مجموعه الصنف",
    ],
    "origin": [
        "origin", "origin type", "origin_type", "local imported", "local/imported", "country of origin",
        "made in", "source", "origine", "origen", "mensei", "herkunft", "المنشا", "بلد المنشا", "محلي مستورد",
    ],
    "min_stock": [
        "min stock", "min_stock", "minimum stock", "minimum", "min qty", "reorder level", "reorder point",
        "safety stock", "stock minimum", "stock minimo", "minimum stok", "mindestbestand",
        "الحد الادني", "حد الطلب", "اقل كميه", "الحد الادني للمخزون",
    ],
    "expiry": [
        "expiry", "expiry date", "expiry_date", "expiration", "expiration date", "exp date", "exp", "expire",
        "expires", "best before", "use by", "date de peremption", "peremption", "date d expiration",
        "fecha de vencimiento", "vencimiento", "caducidad", "fecha de caducidad", "son kullanma tarihi", "skt",
        "verfallsdatum", "mhd", "validade", "data de validade",
        "تاريخ الصلاحيه", "الصلاحيه", "تاريخ الانتهاء", "انتهاء الصلاحيه", "تاريخ انتهاء",
    ],
}

_AR_DIACRITICS = re.compile(r"[ؐ-ًؚ-ٰٟۖ-ۭـ]")
_ARABIC_CHAR = re.compile(r"[؀-ۿ]")
_EASTERN_DIGITS = str.maketrans("٠١٢٣٤٥٦٧٨٩۰۱۲۳۴۵۶۷۸۹٫٬", "01234567890123456789.,")


def normalize_header(s: Any) -> str:
    """Lowercase, strip accents/tashkeel, unify Arabic letter forms, collapse punctuation."""
    if s is None:
        return ""
    t = str(s).strip().lower().translate(_EASTERN_DIGITS)
    t = _AR_DIACRITICS.sub("", t)
    t = (t.replace("أ", "ا").replace("إ", "ا").replace("آ", "ا").replace("ٱ", "ا")
          .replace("ة", "ه").replace("ى", "ي").replace("ؤ", "و").replace("ئ", "ي"))
    # Latin accents: é → e, ç → c, ü → u …
    t = "".join(c for c in unicodedata.normalize("NFKD", t) if not unicodedata.combining(c))
    t = t.replace("ı", "i")
    t = re.sub(r"[%]", " % ", t)
    t = re.sub(r"[^\w%؀-ۿ]+", " ", t)
    t = t.replace("_", " ")
    return re.sub(r"\s+", " ", t).strip()


_NORM_SYNONYMS: dict[str, list[str]] = {
    f: sorted({normalize_header(s) for s in syns}, key=len, reverse=True) for f, syns in SYNONYMS.items()
}


def _header_score(header_norm: str, field: str) -> int:
    """0–100: how well a normalised header matches a field's synonyms."""
    if not header_norm:
        return 0
    best = 0
    h_tokens = header_norm.split()
    padded = f" {header_norm} "
    for syn in _NORM_SYNONYMS[field]:
        if header_norm == syn:
            return 100
        if f" {syn} " in padded:
            # Longer synonym covering more of the header = stronger match.
            coverage = len(syn) / max(len(header_norm), 1)
            best = max(best, int(55 + 35 * coverage))
        elif len(h_tokens) > 1 and all(tok in syn.split() for tok in h_tokens):
            best = max(best, 50)
    return best


# ─── Unit code normalisation ───────────────────────────────────────────────
_UNIT_CODES: dict[str, str] = {}
for _canon, _codes in {
    "box": ["box", "bx", "boxes", "pac", "pack", "pak", "pk", "pkt", "packet", "pck", "bte", "boite", "caja",
            "kutu", "packung", "علبه", "عبوه", "باكو", "علب"],
    "strip": ["strip", "str", "strp", "stp", "strips", "blister", "bls", "blst", "plaquette", "blister pack",
              "شريط", "شرايط", "شرائط", "بلستر"],
    "tablet": ["tablet", "tab", "tabs", "tablets", "tb", "comprime", "comp", "tableta", "tablet", "قرص", "اقراص"],
    "capsule": ["capsule", "cap", "caps", "capsules", "cps", "gelule", "capsula", "kapsul", "كبسوله", "كبسولات"],
    "ampoule": ["ampoule", "amp", "amps", "ampule", "ampul", "ampolla", "امبول", "امبوله", "حقنه"],
    "vial": ["vial", "vl", "vials", "flacon", "فيال", "فيالات"],
    "sachet": ["sachet", "sach", "sac", "sachets", "sobre", "ساشيه", "كيس", "اكياس"],
    "bottle": ["bottle", "btl", "bot", "bottles", "fl", "flask", "frasco", "sise", "flasche", "زجاجه", "ازازه"],
    "tube": ["tube", "tb", "tub", "tubo", "tup", "انبوب", "انبوبه", "تيوب"],
    "piece": ["piece", "pc", "pcs", "pce", "pieces", "ea", "each", "unit", "units", "u", "un", "unite", "unidad",
              "adet", "stuck", "stk", "قطعه", "حبه", "وحده"],
    "carton": ["carton", "ctn", "cartons", "cart", "كرتونه", "كرتون"],
    "jar": ["jar", "pot", "frasco", "برطمان"],
    "can": ["can", "tin", "علبه معدنيه"],
    "kit": ["kit", "طقم"],
    "set": ["set", "مجموعه"],
    "ml": ["ml", "milliliter", "مل"],
}.items():
    for _c in _codes:
        _UNIT_CODES.setdefault(normalize_header(_c), _canon)


def normalize_unit(raw: Any) -> str:
    """Map PAC/STR/TAB/علبة/شريط … to the canonical unit names used by the product form."""
    n = normalize_header(raw)
    if not n:
        return ""
    if n in _UNIT_CODES:
        return _UNIT_CODES[n]
    # "STR 10", "box of 3" → first recognised token
    for tok in n.split():
        if tok in _UNIT_CODES:
            return _UNIT_CODES[tok]
    return n


# ─── Value parsing helpers ─────────────────────────────────────────────────
def parse_number(val: Any) -> Optional[float]:
    """Parse 52, "52.00", "1,234.50", "1.234,50", "٥٢٫٠٠", "EGP 52", "52 ج.م" → float."""
    if val is None or val == "":
        return None
    if isinstance(val, bool):
        return None
    if isinstance(val, (int, float)):
        return float(val)
    s = str(val).strip().translate(_EASTERN_DIGITS)
    s = re.sub(r"[^\d,.\-]", "", s)
    if not s or not re.search(r"\d", s):
        return None
    if "," in s and "." in s:
        # Whichever separator comes last is the decimal point.
        if s.rfind(",") > s.rfind("."):
            s = s.replace(".", "").replace(",", ".")
        else:
            s = s.replace(",", "")
    elif "," in s:
        parts = s.split(",")
        # "1,234" (thousands) vs "12,5" (decimal comma)
        s = s.replace(",", "") if len(parts[-1]) == 3 and len(parts) > 1 else s.replace(",", ".")
    try:
        return float(s)
    except ValueError:
        return None


_DATE_FORMATS = (
    "%Y-%m-%d", "%d/%m/%Y", "%m/%d/%Y", "%d-%m-%Y", "%Y/%m/%d", "%d.%m.%Y", "%Y.%m.%d",
    "%d/%m/%y", "%d-%m-%y", "%m/%y", "%m/%Y", "%m-%Y", "%Y-%m", "%b %Y", "%B %Y", "%d %b %Y", "%d %B %Y",
    "%Y-%m-%d %H:%M:%S", "%d/%m/%Y %H:%M", "%d/%m/%Y %H:%M:%S",
)


def parse_date(val: Any) -> Optional[date]:
    if val in (None, ""):
        return None
    if isinstance(val, datetime):
        return val.date()
    if isinstance(val, date):
        return val
    if isinstance(val, (int, float)) and 20000 < float(val) < 80000:
        # Excel serial date (days since 1899-12-30)
        return (datetime(1899, 12, 30) + timedelta(days=float(val))).date()
    s = str(val).strip().translate(_EASTERN_DIGITS)
    if re.fullmatch(r"\d{5}(\.0+)?", s):
        return parse_date(float(s))
    for fmt in _DATE_FORMATS:
        try:
            d = datetime.strptime(s, fmt).date()
            if fmt in ("%m/%y", "%m/%Y", "%m-%Y", "%Y-%m", "%b %Y", "%B %Y"):
                # Month-only expiry means the end of that month.
                nxt = date(d.year + (d.month // 12), d.month % 12 + 1, 1)
                d = nxt - timedelta(days=1)
            return d
        except ValueError:
            continue
    return None


def _is_gtin(s: str) -> bool:
    if not s.isdigit() or len(s) not in (8, 12, 13, 14):
        return False
    digits = [int(c) for c in s]
    check = digits.pop()
    total = sum(d * (3 if i % 2 == 0 else 1) for i, d in enumerate(reversed(digits)))
    return (10 - total % 10) % 10 == check


def _cell_text(v: Any) -> str:
    if v is None:
        return ""
    if isinstance(v, float) and v.is_integer():
        return str(int(v))
    return str(v).strip()


# ─── File readers → list of sheets, each a list of row lists ───────────────
class _HTMLTableParser(HTMLParser):
    def __init__(self):
        super().__init__()
        self.tables: list[list[list[str]]] = []
        self._row: Optional[list[str]] = None
        self._cell: Optional[list[str]] = None

    def handle_starttag(self, tag, attrs):
        if tag == "table":
            self.tables.append([])
        elif tag == "tr" and self.tables:
            self._row = []
        elif tag in ("td", "th") and self._row is not None:
            self._cell = []

    def handle_endtag(self, tag):
        if tag in ("td", "th") and self._cell is not None and self._row is not None:
            self._row.append(" ".join("".join(self._cell).split()))
            self._cell = None
        elif tag == "tr" and self._row is not None and self.tables:
            self.tables[-1].append(self._row)
            self._row = None

    def handle_data(self, data):
        if self._cell is not None:
            self._cell.append(data)


def _decode_text(content: bytes) -> str:
    if content.startswith((b"\xff\xfe", b"\xfe\xff")):
        return content.decode("utf-16")
    for enc in ("utf-8-sig", "cp1256", "latin-1"):
        try:
            text = content.decode(enc)
            # cp1256 decodes almost anything; only accept it if Arabic letters appear.
            if enc == "cp1256" and not _ARABIC_CHAR.search(text):
                continue
            return text
        except UnicodeDecodeError:
            continue
    return content.decode("utf-8", errors="replace")


def _read_csv(content: bytes) -> list[list[Any]]:
    text = _decode_text(content)
    sample = text[:20000]
    try:
        dialect = csv.Sniffer().sniff(sample, delimiters=",;\t|")
        delim = dialect.delimiter
    except csv.Error:
        counts = {d: sample.count(d) for d in (",", ";", "\t", "|")}
        delim = max(counts, key=counts.get)
    return [row for row in csv.reader(io.StringIO(text), delimiter=delim)]


def read_sheets(content: bytes, filename: str) -> list[tuple[str, list[list[Any]]]]:
    """Return [(sheet_name, rows)] for any supported spreadsheet format."""
    name = (filename or "").lower()
    head = content[:2048].lstrip().lower()

    if content[:4] == b"PK\x03\x04" or name.endswith((".xlsx", ".xlsm")):
        from openpyxl import load_workbook
        wb = load_workbook(io.BytesIO(content), read_only=True, data_only=True)
        try:
            return [(ws.title, [list(r) for r in ws.iter_rows(values_only=True)]) for ws in wb.worksheets]
        finally:
            wb.close()

    if content[:8] == b"\xd0\xcf\x11\xe0\xa1\xb1\x1a\xe1":  # legacy .xls (BIFF)
        try:
            import xlrd
        except ImportError as e:
            raise ValueError("Old .xls files need the 'xlrd' package on the server — "
                             "save the file as .xlsx or CSV and upload again.") from e
        book = xlrd.open_workbook(file_contents=content)
        sheets = []
        for sh in book.sheets():
            rows = []
            for r in range(sh.nrows):
                row = []
                for c in range(sh.ncols):
                    cell = sh.cell(r, c)
                    if cell.ctype == xlrd.XL_CELL_DATE:
                        try:
                            row.append(xlrd.xldate.xldate_as_datetime(cell.value, book.datemode))
                            continue
                        except Exception:
                            pass
                    row.append(cell.value)
                rows.append(row)
            sheets.append((sh.name, rows))
        return sheets

    if head.startswith((b"<!doctype html", b"<html", b"<table", b"<?xml")) or b"<table" in head:
        # Many web/open-source systems "export to Excel" as an HTML table named .xls.
        parser = _HTMLTableParser()
        parser.feed(_decode_text(content))
        tables = [t for t in parser.tables if t]
        if not tables:
            raise ValueError("No table found in the uploaded file")
        return [(f"Table {i + 1}", t) for i, t in enumerate(tables)]

    return [("CSV", _read_csv(content))]


# ─── Header + mapping detection ────────────────────────────────────────────
def _row_header_score(row: list[Any]) -> int:
    score = 0
    for cell in row:
        h = normalize_header(cell)
        if not h or parse_number(cell) is not None and not _ARABIC_CHAR.search(h) and not re.search(r"[a-z]", h):
            continue
        if max(_header_score(h, f) for f in SYNONYMS) >= 60:
            score += 1
    return score


def find_header(sheets: list[tuple[str, list[list[Any]]]]) -> tuple[int, int]:
    """Pick (sheet_index, header_row_index) with the most recognisable headers."""
    best = (0, 0, -1, 0)  # sheet, row, score, data rows
    for si, (_, rows) in enumerate(sheets):
        for ri, row in enumerate(rows[:30]):
            sc = _row_header_score(row)
            data_rows = len(rows) - ri - 1
            if sc > best[2] or (sc == best[2] and sc > 0 and data_rows > best[3]):
                best = (si, ri, sc, data_rows)
    if best[2] <= 0:
        # Nothing recognisable: first non-empty row of the largest sheet.
        si = max(range(len(sheets)), key=lambda i: len(sheets[i][1])) if sheets else 0
        rows = sheets[si][1] if sheets else []
        ri = next((i for i, r in enumerate(rows) if any(_cell_text(c) for c in r)), 0)
        return si, ri
    return best[0], best[1]


def _column_profile(values: list[Any]) -> dict:
    vals = [v for v in values if _cell_text(v) != ""]
    n = len(vals) or 1
    texts = [_cell_text(v) for v in vals]
    nums = [parse_number(v) for v in vals]
    return {
        "filled": len(vals),
        "num_ratio": sum(1 for t, x in zip(texts, nums)
                         if x is not None and not re.search(r"[a-zA-Z؀-ۿ]{2,}", t)) / n,
        # Barcode-shaped: 8–14 digits (valid GTIN checksum counts double).
        "gtin": min(1.0, sum((2 if _is_gtin(t) else 1) for t in texts
                             if t.isdigit() and 8 <= len(t) <= 14) / n),
        "arabic": sum(1 for t in texts if _ARABIC_CHAR.search(t)) / n,
        "date": sum(1 for v in vals if isinstance(v, (date, datetime))
                    or (isinstance(v, str) and parse_date(v) is not None)) / n,
        "avg_len": sum(len(t) for t in texts) / n,
        "unit_like": sum(1 for t in texts if normalize_header(t) in _UNIT_CODES) / n,
        "distinct": len(set(texts)),
    }


def detect_mapping(headers: list[Any], data_rows: list[list[Any]]) -> list[dict]:
    """Return one entry per column: {index, header, field, confidence, samples}."""
    ncols = max([len(headers)] + [len(r) for r in data_rows[:200]] or [0])
    sample = data_rows[:300]
    cols = []
    for ci in range(ncols):
        h = headers[ci] if ci < len(headers) else None
        values = [r[ci] if ci < len(r) else None for r in sample]
        cols.append({
            "index": ci,
            "header": _cell_text(h),
            "norm": normalize_header(h),
            "profile": _column_profile(values),
            "samples": [_cell_text(v) for v in values if _cell_text(v)][:4],
        })

    # 1) Header-name scores, assigned greedily (best pairs first, one field per column).
    pairs = []
    for c in cols:
        if c["profile"]["filled"] == 0 and not c["norm"]:
            continue
        for f in SYNONYMS:
            sc = _header_score(c["norm"], f)
            if sc >= 55:
                pairs.append((sc, -c["index"], c["index"], f))
    pairs.sort(reverse=True)
    assigned: dict[int, tuple[str, int]] = {}
    used: set[str] = set()
    for sc, _, ci, f in pairs:
        if ci in assigned or f in used:
            continue
        assigned[ci] = (f, sc)
        used.add(f)

    # 2) "Code" + "Barcode" both present → code is ours, barcode is the international one.
    generic = next((ci for ci, (f, _) in assigned.items() if f == "barcode_generic"), None)
    if generic is not None:
        sc = assigned[generic][1]
        if "barcode" in used and "intl_barcode" not in used:
            assigned[generic] = ("intl_barcode", sc); used.add("intl_barcode")
        elif "barcode" not in used:
            prof = cols[generic]["profile"]
            target = "intl_barcode" if prof["gtin"] >= 0.6 else "barcode"
            assigned[generic] = (target, sc); used.add(target)
        else:
            del assigned[generic]
        used.discard("barcode_generic")

    # 3) Two name columns: the one written in Arabic is the Arabic name.
    #    (A lone Arabic name column stays name_en; the parser copies it to name_ar.)
    name_col = next((ci for ci, (f, _) in assigned.items() if f == "name_en"), None)
    ar_col = next((ci for ci, (f, _) in assigned.items() if f == "name_ar"), None)
    if name_col is not None and ar_col is not None and cols[name_col]["profile"]["arabic"] >= 0.6 \
            and cols[ar_col]["profile"]["arabic"] < 0.3:
        assigned[name_col], assigned[ar_col] = (("name_ar", assigned[name_col][1]),
                                                ("name_en", assigned[ar_col][1]))

    # 4) Content-based guesses for important fields the headers did not reveal.
    free = [c for c in cols if c["index"] not in assigned and c["profile"]["filled"] > 0]

    def take(field: str, pred, key):
        nonlocal free
        cands = [c for c in free if pred(c["profile"])]
        if field in used or not cands:
            return
        c = max(cands, key=lambda c: key(c["profile"]))
        assigned[c["index"]] = (field, 40)
        used.add(field)
        free = [x for x in free if x["index"] != c["index"]]

    take("intl_barcode", lambda p: p["gtin"] >= 0.6, lambda p: p["gtin"])
    take("expiry", lambda p: p["date"] >= 0.6, lambda p: p["date"])
    take("name_en", lambda p: p["num_ratio"] < 0.3 and p["avg_len"] >= 6 and p["distinct"] > 1,
         lambda p: (p["distinct"], p["avg_len"]))
    take("unit", lambda p: p["unit_like"] >= 0.7, lambda p: p["unit_like"])
    numeric_free = [c for c in free if c["profile"]["num_ratio"] >= 0.9]
    if "price" not in used and len(numeric_free) == 1:
        # Only one unexplained numeric column left — most likely the price.
        assigned[numeric_free[0]["index"]] = ("price", 40)
        used.add("price")

    out = []
    for c in cols:
        f, sc = assigned.get(c["index"], (None, 0))
        out.append({
            "index": c["index"],
            "header": c["header"],
            "field": f,
            "confidence": "high" if sc >= 80 else "medium" if sc >= 55 else "guess" if f else "none",
            "samples": c["samples"],
        })
    return out


# ─── Public API ────────────────────────────────────────────────────────────
def analyze_file(content: bytes, filename: str) -> dict:
    """Detect sheet, header row and column mapping (used by the preview endpoint)."""
    sheets = read_sheets(content, filename)
    if not sheets or not any(rows for _, rows in sheets):
        raise ValueError("The file is empty")
    si, hi = find_header(sheets)
    sheet_name, rows = sheets[si]
    headers = rows[hi] if hi < len(rows) else []
    data = [r for r in rows[hi + 1:] if any(_cell_text(c) for c in r)]
    columns = detect_mapping(list(headers), data)
    return {
        "sheet_index": si,
        "sheet_name": sheet_name,
        "sheets": [n for n, _ in sheets],
        "header_row": hi + 1,
        "row_count": len(data),
        "columns": columns,
    }


def load_mapped_rows(content: bytes, filename: str, mapping: Optional[dict] = None,
                     sheet_index: Optional[int] = None, header_row: Optional[int] = None,
                     on_progress=None) -> tuple[list[dict], dict]:
    """Read the file and return rows keyed by canonical headers for _parse_bulk_row.

    ``mapping`` ({column_index: field_key}) overrides auto-detection — sent by the
    preview screen after the user reviews/corrects the detected columns.
    Returns (rows, info) where info carries the mapping that was applied and the
    first spreadsheet row number of the data (for error messages).
    """
    sheets = read_sheets(content, filename)
    if not sheets:
        raise ValueError("The file is empty")
    auto_si, auto_hi = find_header(sheets)
    si = sheet_index if sheet_index is not None and 0 <= sheet_index < len(sheets) else auto_si
    hi = (header_row - 1) if header_row and header_row >= 1 else (auto_hi if si == auto_si else find_header([sheets[si]])[1])
    rows = sheets[si][1]
    headers = rows[hi] if hi < len(rows) else []
    body = rows[hi + 1:]

    if mapping:
        col_field = {int(k): v for k, v in mapping.items() if v and v in FIELDS}
    else:
        detected = detect_mapping(list(headers), [r for r in body if any(_cell_text(c) for c in r)])
        col_field = {c["index"]: c["field"] for c in detected if c["field"] in FIELDS}

    out: list[dict] = []
    for n, r in enumerate(body, start=1):
        if not any(_cell_text(c) for c in r):
            out.append({})  # keep row numbering aligned with the sheet; skipped by caller
            continue
        rec: dict = {}
        for ci, field in col_field.items():
            if ci < len(r) and r[ci] not in (None, ""):
                rec[FIELDS[field][0]] = r[ci]
        out.append(rec)
        if on_progress and n % 2000 == 0:
            on_progress(message=f"Reading row {n:,}…")
    info = {"first_data_row": hi + 2, "mapping": col_field, "sheet": sheets[si][0]}
    return out, info
